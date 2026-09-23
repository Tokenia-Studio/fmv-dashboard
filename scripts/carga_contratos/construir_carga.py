"""Módulo de contratos FMV — carga inicial determinista (bloque 1.3).

Lee el inventario cruzado (Excel del 07/09/2026) y la carpeta de documentos
renombrada por Daniel, y genera en Mantenimientos/carga_inicial/ (fuera del
repositorio: lleva importes y nombres de FMV):

  carga_contratos.sql   INSERT de contratos, grupos, equipos, obligaciones,
                        documentos y tareas. Una transacción; aborta si el
                        módulo ya tiene datos. Lo ejecuta Carlos en Supabase.
  subida_pdf.json       Qué fichero local va a qué ruta del bucket privado
                        (lo usa scripts/carga_contratos/subir_pdf.js).
  carga.json            Las mismas filas en JSON, para el contraste con la POC.
  informe.md            Recuentos, lo que no cuadra y decisiones aplicadas.

Parte de POC CONTRATOS/construir_datos.py (21/09/2026) con las decisiones de
Carlos del 22 y 23/09/2026 (Arquitectura §10-11, memoria del proyecto):
  · Vista propuesta por categoría (sin confirmar: la confirma dirección).
  · Conjuntos de unidades → grupo + una unidad por equipo (P4).
  · 003 y 095: se respeta la fecha de la hoja como fecha anunciada + tarea (P7).
  · Responsables por defecto: Sachi (Equipos y mantenimiento), Erika (Servicios) (P2).
  · Documentos: los renombrados por Daniel, emparejados con el inventario por
    contenido (SHA-256), más los originales que no pasó a la carpeta nueva.
    No se suben el marcado _DUPLICADO ni lo que no sea PDF (tarea).
  · proveedor_corto = bloque Proveedor del nombre de fichero de Daniel.

No inventa nada: lo que el Excel no trae queda NULL y la app lo enseña como
pendiente. Las fechas próximas no se calculan aquí (las calcula la app).

Uso:  python scripts/carga_contratos/construir_carga.py
"""
import datetime as dt
import hashlib
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

REPO = Path(__file__).resolve().parents[2]
MANT = REPO / "Mantenimientos"
INVENTARIO = MANT / "Inventario contratos FMV 2026-09-07.xlsx"
CARPETA_VIEJA = MANT / "Certificado de servicio y contratos proveedores"
CARPETA_NUEVA = MANT / "CONTRATOS MANTENIMIENTO"
SALIDA = MANT / "carga_inicial"
ORIGEN_TAREA = "inventario 07/09/2026"
RESPONSABLE = {"compras_fabrica": "Sachi", "administracion": "Erika"}

avisos = []  # lo que no cuadra: va al informe


def aviso(texto):
    avisos.append(texto)


# ───────────────────────── Utilidades ─────────────────────────

def norm(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, str):
        s = v.strip()
        return None if s in ("", "—", "---") else s
    return v


def num(v):
    v = norm(v)
    if v is None:
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def fecha_y_precision(v):
    """'2026-05-17' → ('2026-05-17','dia'); '2026-05' → ('2026-05-01','mes'); '2026' → ('2026-01-01','año')."""
    v = norm(v)
    if not v:
        return None, None
    s = str(v)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        dt.date.fromisoformat(s)
        return s, "dia"
    if re.fullmatch(r"\d{4}-\d{2}", s):
        return s + "-01", "mes"
    if re.fullmatch(r"\d{4}", s):
        return s + "-01-01", "año"
    aviso(f"Fecha no reconocida, se deja vacía: {s!r}")
    return None, None


def slug_bloque(texto):
    """Igual que slugBloque() de src/utils/contratosVista.js."""
    s = unicodedata.normalize("NFD", str(texto or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9]+", "-", s)
    return s.strip("-")


def componer_nombre(fecha, precision, proveedor, bloque_tipo, objeto, referencia):
    """Igual que componerNombreFichero() de la app: AAAA-MM-DD_Proveedor_Tipo_Objeto[_Ref].pdf"""
    if not fecha:
        f = "0000-00-00"
    elif precision == "año":
        f = fecha[:4] + "-00-00"
    elif precision == "mes":
        f = fecha[:7] + "-00"
    else:
        f = fecha
    obj = "-".join([p for p in slug_bloque(objeto).split("-") if p][:5]) or "Documento"
    partes = [f, slug_bloque(proveedor) or "Proveedor", bloque_tipo, obj]
    if slug_bloque(referencia):
        partes.append(slug_bloque(referencia))
    return "_".join(partes) + ".pdf"


BLOQUE_TIPO = {"contrato": "Contrato", "renovacion": "Contrato", "anexo": "Contrato", "domiciliacion": "Contrato",
               "presupuesto": "Oferta", "certificado": "Certificado", "parte_visita": "Parte",
               "informe_revision": "Informe", "factura": "Factura", "poliza": "Poliza", "manual": "Manual",
               "legalizacion": "Otro", "otro": "Otro", "pedido": "Otro"}
ROL = {"contrato": "origen", "presupuesto": "origen", "renovacion": "origen", "poliza": "origen",
       "anexo": "origen", "domiciliacion": "origen", "factura": "origen", "pedido": "origen",
       "certificado": "cierre", "parte_visita": "cierre", "informe_revision": "cierre"}


def ids_contrato(texto):
    """'C11 / C12' → ['C11','C12'];  'C22-C28' → ['C22',…,'C28']"""
    if not texto:
        return []
    out = []
    for a, b in re.findall(r"([CH]\d{2})\s*-\s*([CH]\d{2})", texto):
        out += [f"{a[0]}{n:02d}" for n in range(int(a[1:]), int(b[1:]) + 1)]
    out += re.findall(r"[CH]\d{2}", texto)
    return list(dict.fromkeys(out))


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ───────────────────────── Inventario ─────────────────────────
wb = load_workbook(INVENTARIO, data_only=True)


def hoja(nombre, claves):
    filas = []
    for r in wb[nombre].iter_rows(min_row=2, values_only=True):
        if all(c is None for c in r):
            continue
        filas.append({k: norm(v) for k, v in zip(claves, r)})
    return filas


contratos_x = hoja("Contratos", ["id", "proveedor", "provBC", "categoria", "objeto", "nave", "referencia",
    "importe", "periodicidad", "importeAnual", "inicio", "fin", "renovacion", "preaviso", "avisarAntes",
    "enCorreo", "importeCorreo", "pdf", "documento", "estado", "observaciones", "cuentaBC", "activoBC", "confirmado"])
sold_x = hoja("Equipos soldadura", ["num", "activoBC", "modelo", "serie", "asignado", "estado", "ultima", "proxima",
    "pedido", "certificado", "obs", "comentario"])
otros_x = hoja("Otros equipos", ["equipo", "uds", "nave", "identificacion", "periodicidad", "quien", "contrato", "obs"])
docs_x = hoja("Documentos", ["carpeta", "fichero", "tipo", "que", "fecha", "referencia", "escaneado", "legible", "contrato", "notas"])
pend_x = hoja("Pendientes Musachi", ["tipo", "pendiente", "contrato", "respuesta"])
venc_x = hoja("Vencimientos 12 meses", ["fecha", "que", "proveedor", "tipo", "importeAnual", "avisarAntes", "accion", "contrato"])

# ───────────────────────── Documentos: emparejado por contenido ─────────────────────────
viejos = {}  # ruta relativa (como en la hoja) → hash
for r, _, fs in os.walk(CARPETA_VIEJA):
    for f in fs:
        p = Path(r) / f
        viejos[p.relative_to(CARPETA_VIEJA).as_posix()] = sha(p)
nuevos_por_hash = defaultdict(list)
todos_nuevos = []
for r, _, fs in os.walk(CARPETA_NUEVA):
    for f in fs:
        p = Path(r) / f
        todos_nuevos.append(p)
        nuevos_por_hash[sha(p)].append(p)

no_subidos = []  # ficheros de la carpeta nueva que no se suben (duplicado, no PDF)
for p in todos_nuevos:
    if "_DUPLICADO" in p.name:
        no_subidos.append((p, "marcado como duplicado por Daniel"))
    elif p.suffix.lower() != ".pdf":
        no_subidos.append((p, f"no es PDF ({p.suffix}); el almacenamiento solo admite PDF por ahora"))
hash_descartado = {sha(p): motivo for p, motivo in no_subidos}  # su original tampoco se sube

# ───────────────────────── Contratos ─────────────────────────
CATS_FABRICA = {"Mantenimiento", "Inspección", "Residuos", "Certificación", "Calibración", "Legalización"}  # = vistaPropuesta() de la app
NO_VIVOS = {"Histórico", "Sustituido", "Terminado", "Puntual"}

contratos, contrato_por_codigo = [], {}
for i, c in enumerate(contratos_x, 1):
    inicio, inicio_p = fecha_y_precision(c["inicio"])
    fin, fin_p = fecha_y_precision(c["fin"])
    pre = num(c["preaviso"])
    obs = [c["observaciones"]]
    if c["confirmado"]:
        obs.append(f"Compras (inventario 07/09/2026): {c['confirmado']}")
    if c["activoBC"]:
        obs.append(f"Código de activo BC indicado en el inventario: {c['activoBC']}")
    fila = {
        "id": i, "codigo": c["id"], "proveedor_nombre": c["proveedor"],
        # Nº de proveedor en formato BC (6 dígitos con ceros), como el maestro y el diario
        "proveedor_codigo": (str(c["provBC"]).strip().zfill(6) if re.fullmatch(r"\d{1,6}", str(c["provBC"]).strip())
                             else c["provBC"]) if c["provBC"] else None,
        "proveedor_corto": None, "categoria": c["categoria"],
        "vista": "compras_fabrica" if c["categoria"] in CATS_FABRICA else "administracion",
        "vista_confirmada": False, "objeto": c["objeto"], "nave": c["nave"], "referencia": c["referencia"],
        "importe": num(c["importe"]), "periodicidad": c["periodicidad"], "importe_anual": num(c["importeAnual"]),
        "importe_declarado": num(c["importeCorreo"]), "inicio": inicio, "inicio_precision": inicio_p,
        "fin": fin, "fin_precision": fin_p,
        "renovacion": c["renovacion"] if c["renovacion"] in ("tácita", "expresa", "no consta") else None,
        "preaviso_dias": int(pre) if pre is not None else None,
        "estado_documental": c["estado"] or "Por confirmar",
        "vivo": c["id"].startswith("C") and c["estado"] not in NO_VIVOS,
        "cuenta_gasto": c["cuentaBC"], "observaciones": "\n".join(x for x in obs if x) or None,
    }
    for k in ("proveedor_nombre", "categoria", "objeto"):
        if not fila[k]:
            aviso(f"{c['id']}: falta {k} (obligatorio); se pone «Pendiente»")
            fila[k] = "Pendiente"
    if c["renovacion"] and fila["renovacion"] is None:
        aviso(f"{c['id']}: renovación «{c['renovacion']}» no reconocida, se deja vacía")
    contratos.append(fila)
    contrato_por_codigo[c["id"]] = fila

# ───────────────────────── Documentos ─────────────────────────
documentos, documento_contrato = [], []
rutas_usadas = set()
subida = []
hash_ya_subido = {}
filas_doc_sin_fichero = 0

originales_pendientes = []


def procesar_documento(d, segunda_pasada=False):
    global filas_doc_sin_fichero
    if not d["fichero"] or d["fichero"].endswith("/"):
        filas_doc_sin_fichero += 1
        return
    h = viejos.get(d["fichero"])
    if h is None:
        aviso(f"Documento de la hoja no está en disco: {d['fichero']}")
        return
    if h in hash_descartado:
        aviso(f"No se sube {d['fichero']}: su copia en la carpeta de Daniel está descartada ({hash_descartado[h]})")
        return
    ids = [x for x in ids_contrato(d["contrato"]) if x in contrato_por_codigo]
    for x in ([] if segunda_pasada else ids_contrato(d["contrato"])):
        if x not in contrato_por_codigo:
            aviso(f"Documento {d['fichero']}: contrato {x} no existe en la hoja Contratos")
    candidatos = [p for p in nuevos_por_hash.get(h, []) if not any(p == q for q, _ in no_subidos)]
    tipo = d["tipo"] or "otro"
    if tipo not in BLOQUE_TIPO:
        aviso(f"{d['fichero']}: tipo «{tipo}» desconocido, se carga como «otro»")
        tipo = "otro"

    if h in hash_ya_subido:
        # Mismo contenido que otro fichero ya cargado: una sola ficha, enlazada a los contratos de ambos
        doc_id = hash_ya_subido[h]
        for cid in ids:
            par = (doc_id, contrato_por_codigo[cid]["id"])
            if par not in documento_contrato:
                documento_contrato.append(par)
        aviso(f"{d['fichero']}: mismo contenido que otro documento ya cargado; se enlaza a esa ficha")
        return

    if not candidatos and not segunda_pasada:
        originales_pendientes.append(d)  # se sube después, con el nombre corto de su contrato
        return

    if candidatos:
        local = candidatos[0]
        nombre = local.name
        partes = nombre.split("_")
        proveedor = partes[1]
        m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", partes[0])
        if partes[0] == "0000-00-00" or not m:
            fecha, fprec = fecha_y_precision(d["fecha"])
        elif m.group(3) == "00" and m.group(2) == "00":
            fecha, fprec = f"{m.group(1)}-01-01", "año"
        elif m.group(3) == "00":
            fecha, fprec = f"{m.group(1)}-{m.group(2)}-01", "mes"
        else:
            fecha, fprec = partes[0], "dia"
        origen_fichero = "renombrado por Daniel"
    else:
        # Original que no está en la carpeta nueva: se sube con el nombre que compone la app
        local = CARPETA_VIEJA / d["fichero"]
        if local.suffix.lower() != ".pdf" or not local.read_bytes()[:5] == b"%PDF-":
            vacio = local.stat().st_size == 0
            aviso(f"No se sube el original {d['fichero']}: {'está vacío (0 bytes)' if vacio else 'no es un PDF'}")
            return
        fecha, fprec = fecha_y_precision(d["fecha"])
        corto = contrato_por_codigo[ids[0]]["proveedor_corto"] if ids else None
        cprov = contrato_por_codigo[ids[0]]["proveedor_nombre"] if ids else d["carpeta"]
        proveedor = corto or slug_bloque(re.split(r"[\s/(,]", cprov.strip())[0])
        nombre = componer_nombre(fecha, fprec, proveedor, BLOQUE_TIPO[tipo], d["que"] or tipo, d["referencia"])
        origen_fichero = "original (no estaba en la carpeta de Daniel)"
        aviso(f"Se sube el original {d['fichero']} como {proveedor}/{nombre}")

    if Path(local).read_bytes()[:5] != b"%PDF-":
        aviso(f"No se sube {Path(local).name}: no es un PDF válido")
        return
    ruta = f"{proveedor}/{nombre}"
    base = ruta
    n = 2
    while ruta in rutas_usadas:
        ruta = base[:-4] + f"-{n}.pdf"
        n += 1
    rutas_usadas.add(ruta)
    doc_id = len(documentos) + 1
    hash_ya_subido[h] = doc_id
    documentos.append({
        "id": doc_id, "ruta": ruta, "nombre_original": nombre, "tipo": tipo, "rol": ROL.get(tipo, "otro"),
        "fecha": fecha, "fecha_precision": fprec, "referencia": d["referencia"], "descripcion": d["que"],
        "legible": (d["legible"] or "").upper() != "NO", "equipo_id": None, "obligacion_id": None,
    })
    subida.append({"local": str(local.relative_to(REPO)), "ruta": ruta, "sha256": h, "origen": origen_fichero})
    for cid in ids:
        documento_contrato.append((doc_id, contrato_por_codigo[cid]["id"]))
    # Nombre corto del proveedor para los contratos del documento (bloque Proveedor de Daniel)
    for cid in ids:
        c = contrato_por_codigo[cid]
        if c["proveedor_corto"] and c["proveedor_corto"] != proveedor:
            c.setdefault("_otros_cortos", set()).add(proveedor)
        elif not c["proveedor_corto"]:
            c["proveedor_corto"] = proveedor

for d in docs_x:
    procesar_documento(d)

for c in contratos:
    if c.get("_otros_cortos"):
        aviso(f"{c['codigo']}: sus documentos usan varios nombres cortos ({c['proveedor_corto']}, "
              f"{', '.join(sorted(c['_otros_cortos']))}); se toma {c['proveedor_corto']}")
    c.pop("_otros_cortos", None)
    if not c["proveedor_corto"]:
        c["proveedor_corto"] = slug_bloque(re.split(r"[\s/(,]", c["proveedor_nombre"].strip())[0]) or None
        c["_corto_propuesto"] = True

for d in originales_pendientes:
    procesar_documento(d, segunda_pasada=True)

# ───────────────────────── Grupos, equipos y obligaciones ─────────────────────────
TIPOS_EQUIPO = [("puente grúa", "Elevación"), ("polipasto", "Elevación"), ("eslinga", "Elevación"),
    ("extintor", "Contra incendios"), ("compresor", "Aire comprimido"), ("secador", "Aire comprimido"),
    ("plegadora", "Maquinaria"), ("carretilla", "Carretillas"), ("split", "Climatización"),
    ("térmica", "Climatización"), ("generador", "Climatización"), ("radial", "Herramienta en renting"),
    ("extracción", "Instalaciones"), ("puerta", "Instalaciones"), ("estanter", "Instalaciones"),
    ("tanque", "Instalaciones")]
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre"]

grupos, equipos, obligaciones, contrato_equipo, tareas = [], [], [], [], []


def nueva_obligacion(**kw):
    base = {"id": len(obligaciones) + 1, "equipo_id": None, "grupo_id": None, "contrato_id": None,
            "tipo": "Revisión", "etiqueta": "Revisión", "periodicidad_meses": None, "proveedor_nombre": None,
            "precio": None, "pedido_pcp": None, "primera_fecha": None, "primera_fecha_precision": None,
            "mes_habitual": None, "fecha_anunciada": None, "fecha_anunciada_precision": None, "activa": True}
    base.update(kw)
    obligaciones.append(base)
    return base


def nueva_tarea(tipo, texto, vista, contrato_id=None, equipo_id=None, respuesta=None, responsable=None):
    tareas.append({"id": len(tareas) + 1, "tipo": tipo, "texto": texto, "respuesta": respuesta,
                   "responsable": responsable or RESPONSABLE[vista], "estado": "Abierta", "fecha_limite": None,
                   "origen": ORIGEN_TAREA if respuesta is not None or tipo != "Confirmar" else ORIGEN_TAREA,
                   "vista": vista, "contrato_id": contrato_id, "equipo_id": equipo_id})


def periodicidad_meses(txt):
    t = (txt or "").lower()
    if re.search(r"sin (mantenimiento|revisión)", t):
        return "sin_plan"
    if re.search(r"semestral|6 meses|2 visitas", t):
        return 6
    if re.search(r"anual|/año", t):
        return 12
    return None


def ultimo_cierre(ids_codigo):
    """Documento de cierre más reciente (con fecha de día) de esos contratos → (fecha, doc_id)."""
    ids = {contrato_por_codigo[x]["id"] for x in ids_codigo if x in contrato_por_codigo}
    cand = [(documentos[doc - 1]["fecha"], doc) for doc, cid in documento_contrato
            if cid in ids and documentos[doc - 1]["rol"] == "cierre" and documentos[doc - 1]["fecha_precision"] == "dia"]
    return max(cand) if cand else (None, None)


def regimen_de(texto, ids_codigo):
    t = (texto or "").lower()
    cats = {contrato_por_codigo[x]["categoria"] for x in ids_codigo if x in contrato_por_codigo}
    if "alquiler" in t:
        return "alquiler"
    if "renting" in t or "Renting" in cats:
        return "renting"
    return "propio"


def enlazar(equipo_id, ids_codigo):
    for x in ids_codigo:
        if x in contrato_por_codigo:
            par = (contrato_por_codigo[x]["id"], equipo_id)
            if par not in contrato_equipo:
                contrato_equipo.append(par)


# Grupos de soldadura: un equipo por grupo, una calibración anual por equipo (contrato C49)
C49 = "C49" if "C49" in contrato_por_codigo else None
for s in sold_x:
    eid = len(equipos) + 1
    estado = s["estado"] if s["estado"] in ("activo", "baja", "cedido", "en reparación") else "activo"
    if s["estado"] and s["estado"] != estado:
        pass  # «sin calibración registrada»: equipo activo sin fecha → la app genera la tarea
    obs = [x for x in (s["obs"], s["comentario"]) if x]
    equipos.append({"id": eid, "grupo_id": None, "nombre": f"Grupo de soldadura {s['num']}" if s["num"] else f"Grupo de soldadura (sin nº) {s['modelo'] or ''}".strip(),
        "tipo": "Grupo de soldadura", "num_interno": s["num"], "modelo": s["modelo"], "num_serie": s["serie"],
        "identificacion": None, "nave": "Gavilanes 21", "asignado_a": s["asignado"], "estado": estado,
        "regimen": "propio", "activo_fijo_bc": s["activoBC"], "sin_plan": False,
        "observaciones": " · ".join(obs) or None})
    if C49:
        enlazar(eid, [C49])
    if estado == "baja":
        continue
    ultima, ultima_p = fecha_y_precision(s["ultima"])
    anunciada, anunciada_p = None, None
    if ultima and s["proxima"] and re.fullmatch(r"\d{4}-\d{2}", str(s["proxima"])):
        esperado = f"{int(ultima[:4]) + 1}-{ultima[5:7]}"
        if s["proxima"] != esperado and s["proxima"] > ultima[:7]:
            anunciada, anunciada_p = s["proxima"] + "-01", "mes"
            nueva_tarea("Confirmar", f"{equipos[-1]['nombre']}: la hoja F16-C da como próxima calibración {s['proxima']}, "
                        f"que no es última + 12 meses ({esperado}). Confirmar la periodicidad (P7).",
                        "compras_fabrica", equipo_id=eid)
    nueva_obligacion(equipo_id=eid, contrato_id=contrato_por_codigo[C49]["id"] if C49 else None,
        proveedor_nombre=contrato_por_codigo[C49]["proveedor_nombre"] if C49 else None, tipo="Calibración",
        etiqueta="Calibración anual", periodicidad_meses=12, pedido_pcp=s["pedido"],
        primera_fecha=ultima, primera_fecha_precision=ultima_p,
        fecha_anunciada=anunciada, fecha_anunciada_precision=anunciada_p)

# Otros equipos
per_anterior = None
for o in otros_x:
    if (o["periodicidad"] or "").strip().lower() in ("ídem", "idem"):
        o["periodicidad"] = per_anterior
    per_anterior = o["periodicidad"]
    if o["equipo"].startswith("Grupos de soldadura"):
        continue  # fila resumen: el detalle ya está equipo a equipo
    ids = ids_contrato(o["contrato"])
    for x in ids:
        if x not in contrato_por_codigo:
            aviso(f"Equipo «{o['equipo']}»: contrato {x} no existe en la hoja Contratos")
    ids = [x for x in ids if x in contrato_por_codigo]
    cats = {contrato_por_codigo[x]["categoria"] for x in ids}
    per = periodicidad_meses(o["periodicidad"])
    es_inspeccion = "inspecci" in (o["periodicidad"] or "").lower()
    fecha_cierre, _ = ultimo_cierre(ids)
    pista, pista_p = None, None
    m = re.search(r"(\d{2})/(\d{2})/(\d{4})", o["periodicidad"] or "")
    if m:
        pista, pista_p = f"{m.group(3)}-{m.group(2)}-{m.group(1)}", "dia"
    m = re.search(r"Próxima:\s*(\w+)\s+(\d{4})", o["obs"] or "")
    if m and m.group(1).lower() in MESES:
        pista, pista_p = f"{m.group(2)}-{MESES.index(m.group(1).lower()) + 1:02d}-01", "mes"
    paren = re.search(r"\(([^)]*)\)", o["periodicidad"] or "")
    meses = [MESES.index(w) + 1 for w in re.findall(r"[a-záéíóú]+", paren.group(1).lower()) if w in MESES] if paren else []
    comunes = dict(contrato_id=contrato_por_codigo[ids[0]]["id"] if ids else None, proveedor_nombre=o["quien"],
                   tipo="Inspección" if es_inspeccion else "Revisión",
                   periodicidad_meses=per if isinstance(per, int) else None,
                   primera_fecha=fecha_cierre, primera_fecha_precision="dia" if fecha_cierre else None)

    if cats == {"Seguridad"}:
        # Intrusión y alarmas: la revisión es obligación del contrato, no de un equipo
        nueva_obligacion(**comunes, etiqueta=f"Revisión anual · {o['equipo']}", fecha_anunciada=pista,
                         fecha_anunciada_precision=pista_p)
        continue

    uds = num(o["uds"])
    tipo = next((t for k, t in TIPOS_EQUIPO if k in o["equipo"].lower()), "Instalaciones")
    regimen = regimen_de(o["equipo"], ids)
    obs = [x for x in (o["obs"], f"Periodicidad según inventario: {o['periodicidad']}" if o["periodicidad"] else None) if x]
    sin_plan = per == "sin_plan"

    if uds is None or uds > 1:
        # Conjunto de unidades → grupo + una unidad por elemento (P4)
        gid = len(grupos) + 1
        grupos.append({"id": gid, "nombre": o["equipo"], "tipo": tipo, "nave": o["nave"],
                       "observaciones": " · ".join(x for x in (o["identificacion"] and f"Identificación: {o['identificacion']}", *obs) if x) or None})
        n = int(uds) if uds else 0
        base = re.split(r"[(+]", o["equipo"])[0].strip()
        for k in range(1, n + 1):
            eid = len(equipos) + 1
            equipos.append({"id": eid, "grupo_id": gid, "nombre": f"{base} {k:02d}", "tipo": tipo, "num_interno": None,
                "modelo": None, "num_serie": None, "identificacion": None, "nave": o["nave"], "asignado_a": None,
                "estado": "activo", "regimen": regimen, "activo_fijo_bc": None, "sin_plan": sin_plan, "observaciones": None})
            enlazar(eid, ids)
        if not n:
            nueva_tarea("Confirmar", f"«{o['equipo']}»: el inventario no indica cuántas unidades hay. Añadirlas al grupo desde la app.",
                        "compras_fabrica")
        m = re.search(r"(\d+) NO APTOS", o["obs"] or "")
        if m:
            nueva_tarea("Confirmar", f"«{o['equipo']}»: {m.group(1)} de {n} no aptos en la última inspección. "
                        "Identificar cuáles al registrar la próxima (resultado unidad a unidad).", "compras_fabrica")
        sujeto = {"grupo_id": gid}
    else:
        eid = len(equipos) + 1
        equipos.append({"id": eid, "grupo_id": None, "nombre": o["equipo"], "tipo": tipo, "num_interno": None,
            "modelo": None, "num_serie": None, "identificacion": o["identificacion"], "nave": o["nave"], "asignado_a": None,
            "estado": "activo", "regimen": regimen, "activo_fijo_bc": None, "sin_plan": sin_plan,
            "observaciones": " · ".join(obs) or None})
        enlazar(eid, ids)
        sujeto = {"equipo_id": eid}
        m = re.search(r"(\d+) NO APTOS", o["obs"] or "")
        if m:
            aviso(f"«{o['equipo']}»: «{m.group(0)}» en un equipo de una unidad; revisar a mano")

    if sin_plan:
        continue
    if len(meses) > 1:
        for mes in meses:
            nueva_obligacion(**sujeto, **comunes, etiqueta=f"Revisión anual ({MESES[mes - 1]})", mes_habitual=mes)
    else:
        nueva_obligacion(**sujeto, **comunes, etiqueta=(o["periodicidad"] or "Revisión")[:120],
                         mes_habitual=meses[0] if meses else None, fecha_anunciada=pista, fecha_anunciada_precision=pista_p)

# ───────────────────────── Tareas del inventario (hoja Pendientes) ─────────────────────────
TIPOS_TAREA = {"Falta documento", "Discrepancia", "Confirmar", "Decidir", "Otro"}
tareas_inventario = []
for p in pend_x:
    ids = [x for x in ids_contrato(p["contrato"]) if x in contrato_por_codigo]
    c = contrato_por_codigo[ids[0]] if ids else None
    vista = c["vista"] if c else "compras_fabrica"
    tipo = p["tipo"] if p["tipo"] in TIPOS_TAREA else "Otro"
    texto = p["pendiente"] if tipo == p["tipo"] else f"[{p['tipo']}] {p['pendiente']}"
    if len(ids) > 1:
        texto += f" (contratos {', '.join(ids)})"
    tareas_inventario.append(dict(tipo=tipo, texto=texto, vista=vista, contrato_id=c["id"] if c else None,
                                  respuesta=p["respuesta"], responsable=p["respuesta"] if p["respuesta"] == "Erika" else None))
# Las del inventario primero (conservan el orden de la hoja), después las que añade la carga
extra = tareas[:]
tareas.clear()
for t in tareas_inventario:
    nueva_tarea(**t)
for t in extra:
    t["id"] = len(tareas) + 1
    tareas.append(t)
for p, motivo in no_subidos:
    if p.suffix.lower() != ".pdf":
        corto = p.name.split("_")[1] if "_" in p.name else None
        c = next((c for c in contratos if c["proveedor_corto"] == corto), None)
        nueva_tarea("Falta documento", f"No se ha cargado «{p.name}»: {motivo}. Guardarlo como PDF y subirlo a su contrato.",
                    c["vista"] if c else "compras_fabrica", contrato_id=c["id"] if c else None)

# ───────────────────────── Hitos sueltos (US-012) ─────────────────────────
# Filas de la hoja «Vencimientos 12 meses» que solo existen como texto: un pago
# aplazado o una revisión de precios. No son vencimiento del contrato ni revisión
# periódica, así que el calendario no las genera solo: se cargan como hito.
hitos = []
for v in venc_x:
    que = (v["que"] or "").lower()
    if not v["contrato"] or v["contrato"] not in contrato_por_codigo:
        continue
    if "pago" in que:
        tipo_h = "Pago"
    elif "revisión de precios" in que or "revision de precios" in que:
        tipo_h = "Renegociación"
    else:
        continue
    fecha_h, _ = fecha_y_precision(str(v["fecha"])[:10])
    if not fecha_h:
        continue
    m = re.search(r"preaviso\s+(\d+)\s*d", que)
    hitos.append({"id": len(hitos) + 1, "contrato_id": contrato_por_codigo[v["contrato"]]["id"], "tipo": tipo_h,
                  "descripcion": v["que"], "fecha": fecha_h, "importe": None,
                  "aviso_dias": int(m.group(1)) if m else 30, "cerrado": False})
    aviso(f"Hito cargado desde la hoja de vencimientos: {v['contrato']} {fecha_h} — {v['que']}")

# ───────────────────────── SQL ─────────────────────────

def lit(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(round(v, 2)) if isinstance(v, float) else str(v)
    return "'" + str(v).replace("'", "''") + "'"


def inserts(tabla, filas, columnas):
    if not filas:
        return f"-- {tabla}: sin filas\n"
    lineas = [f"insert into public.{tabla} ({', '.join(columnas)}) values"]
    lineas += [("  (" + ", ".join(lit(f[c]) for c in columnas) + ")") + ("," if i < len(filas) - 1 else ";")
               for i, f in enumerate(filas)]
    return "\n".join(lineas) + "\n"


COL = {
    "ctr_contratos": ["id", "codigo", "proveedor_nombre", "proveedor_codigo", "proveedor_corto", "categoria", "vista",
                      "vista_confirmada", "objeto", "nave", "referencia", "importe", "periodicidad", "importe_anual",
                      "importe_declarado", "inicio", "inicio_precision", "fin", "fin_precision", "renovacion",
                      "preaviso_dias", "estado_documental", "vivo", "cuenta_gasto", "observaciones"],
    "ctr_grupos_equipos": ["id", "nombre", "tipo", "nave", "observaciones"],
    "ctr_equipos": ["id", "grupo_id", "nombre", "tipo", "num_interno", "modelo", "num_serie", "identificacion", "nave",
                    "asignado_a", "estado", "regimen", "activo_fijo_bc", "sin_plan", "observaciones"],
    "ctr_obligaciones": ["id", "equipo_id", "grupo_id", "contrato_id", "tipo", "etiqueta", "periodicidad_meses",
                         "proveedor_nombre", "precio", "pedido_pcp", "primera_fecha", "primera_fecha_precision",
                         "mes_habitual", "fecha_anunciada", "fecha_anunciada_precision", "activa"],
    "ctr_documentos": ["id", "ruta", "nombre_original", "tipo", "rol", "fecha", "fecha_precision", "referencia",
                       "descripcion", "legible", "equipo_id", "obligacion_id"],
    "ctr_tareas": ["id", "tipo", "texto", "respuesta", "responsable", "estado", "fecha_limite", "origen", "vista",
                   "contrato_id", "equipo_id"],
    "ctr_hitos": ["id", "contrato_id", "tipo", "descripcion", "fecha", "importe", "aviso_dias", "cerrado"],
}
TABLAS_CON_ID = ["ctr_contratos", "ctr_grupos_equipos", "ctr_equipos", "ctr_obligaciones", "ctr_documentos",
                 "ctr_tareas", "ctr_hitos"]
TABLAS_MODULO = ["ctr_contratos", "ctr_grupos_equipos", "ctr_equipos", "ctr_contrato_equipo", "ctr_obligaciones",
                 "ctr_realizadas", "ctr_documentos", "ctr_documento_contrato", "ctr_tareas", "ctr_tareas_notas"]
filas_sql = {"ctr_contratos": contratos, "ctr_grupos_equipos": grupos, "ctr_equipos": equipos,
             "ctr_obligaciones": obligaciones, "ctr_documentos": documentos, "ctr_tareas": tareas, "ctr_hitos": hitos}
esperado = {**{t: len(f) for t, f in filas_sql.items()},
            "ctr_contrato_equipo": len(contrato_equipo), "ctr_documento_contrato": len(documento_contrato)}

sql = [f"""-- ============================================================================
-- MÓDULO DE CONTRATOS · CARGA INICIAL (bloque 1.3)
-- Generado el {dt.date.today().isoformat()} por scripts/carga_contratos/construir_carga.py
-- Fuente: {INVENTARIO.name} + carpeta «{CARPETA_NUEVA.name}» (renombrado de Daniel)
-- ============================================================================
-- CONFIDENCIAL: lleva importes y proveedores de FMV. No va al repositorio.
--
-- Orden: 1) sql/contratos_proveedor_corto_2026-09-23.sql
--        2) node scripts/carga_contratos/subir_pdf.js        (sube los PDF)
--        3) ESTE fichero (editor SQL de Supabase: Ctrl+A, Ctrl+V, Ctrl+Enter)
-- Se ejecuta UNA vez: aborta sin tocar nada si el módulo ya tiene datos.
-- Una sola transacción: o entra todo o no entra nada.
-- ============================================================================
begin;

do $g$
declare n bigint;
begin
  if not exists (select 1 from information_schema.columns where table_schema = 'public'
                 and table_name = 'ctr_contratos' and column_name = 'proveedor_corto') then
    raise exception 'Falta la columna proveedor_corto: ejecutar antes sql/contratos_proveedor_corto_2026-09-23.sql';
  end if;
  select (select count(*) from public.ctr_contratos) + (select count(*) from public.ctr_equipos)
       + (select count(*) from public.ctr_grupos_equipos) + (select count(*) from public.ctr_obligaciones)
       + (select count(*) from public.ctr_documentos) + (select count(*) from public.ctr_tareas)
       + (select count(*) from public.ctr_hitos) into n;
  if n > 0 then
    raise exception 'El módulo ya tiene % filas: la carga inicial solo se hace sobre tablas vacías. No se ha tocado nada.', n;
  end if;
end;
$g$;
"""]
for t in TABLAS_CON_ID:
    sql.append(f"\n-- {t}: {len(filas_sql[t])} filas\n" + inserts(t, filas_sql[t], COL[t]))
sql.append(f"\n-- ctr_contrato_equipo: {len(contrato_equipo)} enlaces\n" + inserts(
    "ctr_contrato_equipo", [{"contrato_id": a, "equipo_id": b} for a, b in contrato_equipo], ["contrato_id", "equipo_id"]))
sql.append(f"\n-- ctr_documento_contrato: {len(documento_contrato)} enlaces\n" + inserts(
    "ctr_documento_contrato", [{"documento_id": a, "contrato_id": b} for a, b in documento_contrato], ["documento_id", "contrato_id"]))
sql.append("\n-- Secuencias: que los nuevos registros de la app sigan tras los cargados\n")
for t in TABLAS_CON_ID:
    sql.append(f"select setval(pg_get_serial_sequence('public.{t}', 'id'), greatest((select max(id) from public.{t}), 1));\n")
comprob = " ".join(
    f"if (select count(*) from public.{t}) <> {n} then raise exception 'Carga incompleta en {t}: se esperaban {n} filas'; end if;"
    for t, n in esperado.items())
sql.append(f"""
do $c$ begin {comprob} end; $c$;

commit;

-- Resultado
select 'contratos' as tabla, (select count(*) from public.ctr_contratos) as filas, {esperado['ctr_contratos']} as esperado
union all select 'grupos de equipos', (select count(*) from public.ctr_grupos_equipos), {esperado['ctr_grupos_equipos']}
union all select 'equipos (unidades incluidas)', (select count(*) from public.ctr_equipos), {esperado['ctr_equipos']}
union all select 'obligaciones', (select count(*) from public.ctr_obligaciones), {esperado['ctr_obligaciones']}
union all select 'documentos', (select count(*) from public.ctr_documentos), {esperado['ctr_documentos']}
union all select 'tareas', (select count(*) from public.ctr_tareas), {esperado['ctr_tareas']}
union all select 'hitos', (select count(*) from public.ctr_hitos), {esperado['ctr_hitos']}
union all select 'enlaces contrato-equipo', (select count(*) from public.ctr_contrato_equipo), {esperado['ctr_contrato_equipo']}
union all select 'enlaces documento-contrato', (select count(*) from public.ctr_documento_contrato), {esperado['ctr_documento_contrato']}
union all select 'documentos cuyo PDF no está en el bucket',
       (select count(*) from public.ctr_documentos d
        where not exists (select 1 from storage.objects o where o.bucket_id = 'contratos' and o.name = d.ruta)), 0;
""")

# ───────────────────────── Salida ─────────────────────────
SALIDA.mkdir(exist_ok=True)
(SALIDA / "carga_contratos.sql").write_text("".join(sql), encoding="utf-8")
(SALIDA / "subida_pdf.json").write_text(json.dumps(subida, ensure_ascii=False, indent=1), encoding="utf-8")
(SALIDA / "carga.json").write_text(json.dumps({
    "contratos": contratos, "grupos": grupos, "equipos": equipos, "obligaciones": obligaciones,
    "documentos": documentos, "tareas": tareas, "hitos": hitos,
    "contratoEquipo": [{"contrato_id": a, "equipo_id": b} for a, b in contrato_equipo],
    "documentoContrato": [{"documento_id": a, "contrato_id": b} for a, b in documento_contrato],
    "vencimientosHoja": venc_x}, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

vivos = [c for c in contratos if c["vivo"]]
inf = [f"# Carga inicial del módulo de contratos — informe\n",
       f"Generado el {dt.datetime.now():%d/%m/%Y %H:%M} · fuente `{INVENTARIO.name}` y carpeta `{CARPETA_NUEVA.name}`\n",
       "## Recuentos\n", "| Qué | Filas |", "|---|---|"]
inf += [f"| {t} | {n} |" for t, n in esperado.items()]
inf += ["", "## Contratos por vista propuesta (sin confirmar)\n",
        f"- Equipos y mantenimiento: {sum(1 for c in contratos if c['vista'] == 'compras_fabrica')} "
        f"({sum(1 for c in vivos if c['vista'] == 'compras_fabrica')} vivos)",
        f"- Servicios y arrendamientos: {sum(1 for c in contratos if c['vista'] == 'administracion')} "
        f"({sum(1 for c in vivos if c['vista'] == 'administracion')} vivos)",
        f"- Importe anual vivo: {sum(c['importe_anual'] or 0 for c in vivos):,.2f} € "
        f"(vigente {sum(c['importe_anual'] or 0 for c in vivos if c['estado_documental'] == 'Vigente'):,.2f} €)".replace(",", "X").replace(".", ",").replace("X", "."),
        "", "## Documentos\n",
        f"- En la carpeta de Daniel: {len(todos_nuevos)} ficheros. Se suben {len(subida)} PDF "
        f"({sum(1 for s in subida if s['origen'].startswith('renombrado'))} renombrados + "
        f"{sum(1 for s in subida if s['origen'].startswith('original'))} original(es) que no estaba(n)).",
        f"- Filas de la hoja Documentos sin fichero (carpetas vacías): {filas_doc_sin_fichero}.",
        "- No se suben:"]
inf += [f"  - `{p.name}`: {m}" for p, m in no_subidos]
inf += ["", "## Nombre corto de proveedor\n",
        "Del bloque Proveedor de los ficheros de Daniel. Propuestos (sin documento del que sacarlo): "
        + (", ".join(f"{c['codigo']} → {c['proveedor_corto']}" for c in contratos if c.get("_corto_propuesto")) or "ninguno")]
inf += ["", "## Equipos\n",
        f"- Grupos: {len(grupos)} (" + ", ".join(f"{g['nombre']}: {sum(1 for e in equipos if e['grupo_id'] == g['id'])} uds" for g in grupos) + ")",
        f"- Régimen renting/alquiler: " + (", ".join(sorted({e['nombre'] if not e['grupo_id'] else next(g['nombre'] for g in grupos if g['id'] == e['grupo_id']) for e in equipos if e['regimen'] != 'propio'})) or "ninguno"),
        f"- Sin plan de mantenimiento: " + (", ".join(sorted({e['nombre'] if not e['grupo_id'] else next(g['nombre'] for g in grupos if g['id'] == e['grupo_id']) for e in equipos if e['sin_plan']})) or "ninguno"),
        f"- Obligaciones sin fecha de partida (saldrán «sin fecha fijada»): {sum(1 for o in obligaciones if not o['primera_fecha'] and not o['fecha_anunciada'] and not o['mes_habitual'])}",
        "", "## Tareas\n",
        f"- Del inventario: {len(tareas_inventario)} · añadidas por la carga: {len(tareas) - len(tareas_inventario)}",
        f"- Responsable por defecto: Sachi {sum(1 for t in tareas if t['responsable'] == 'Sachi')} · Erika {sum(1 for t in tareas if t['responsable'] == 'Erika')}",
        "", f"## Avisos ({len(avisos)})\n"]
inf += [f"- {a}" for a in avisos] or ["- Ninguno"]
(SALIDA / "informe.md").write_text("\n".join(inf) + "\n", encoding="utf-8")

print("\n".join(inf))
print(f"\nEscrito en {SALIDA.relative_to(REPO)}: carga_contratos.sql, subida_pdf.json, carga.json, informe.md")
