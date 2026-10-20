# -*- coding: utf-8 -*-
"""
Manual de usuario — Módulo Presupuesto (PyG Presupuesto vs Real) del FMV Dashboard.
TOKENIA Automation SLU.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AMBER=RGBColor(0xB4,0x5A,0x00); DARK=RGBColor(0x1F,0x29,0x37)
GREY=RGBColor(0x55,0x5B,0x66); LIGHT=RGBColor(0x8A,0x8F,0x98); WHITE=RGBColor(0xFF,0xFF,0xFF)
RED=RGBColor(0xB1,0x2A,0x1B); GREEN=RGBColor(0x1E,0x7A,0x34)
HEADBG="1F2937"; ZEBRA="F3F4F6"; AMBERBG="FEF3C7"; REDBG="FDE8E6"; GREENBG="E8F4EA"; BLUEBG="E8EFF8"

doc=Document()
nn=doc.styles['Normal']; nn.font.name='Calibri'; nn.font.size=Pt(10.5); nn.font.color.rgb=DARK
for s in doc.sections:
    s.top_margin=Inches(0.9); s.bottom_margin=Inches(0.9); s.left_margin=Inches(0.9); s.right_margin=Inches(0.9)

def shade(c,h):
    p=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),h); p.append(sh)
def sct(c,t,bold=False,color=DARK,size=10,align='left',white=False):
    c.text=''; p=c.paragraphs[0]
    p.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r=p.add_run(t); r.bold=bold; r.font.size=Pt(size); r.font.color.rgb=WHITE if white else color
    p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2)
def h1(t,num=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(16); p.paragraph_format.space_after=Pt(4)
    r=p.add_run((f'{num}.  ' if num else '')+t); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=AMBER
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); bo=OxmlElement('w:bottom')
    bo.set(qn('w:val'),'single'); bo.set(qn('w:sz'),'6'); bo.set(qn('w:space'),'4'); bo.set(qn('w:color'),'E0A458')
    b.append(bo); pPr.append(b)
def h2(t):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(t); r.bold=True; r.font.size=Pt(11.5); r.font.color.rgb=DARK
def body(t,size=10.5,color=DARK,space=4):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(space)
    r=p.add_run(t); r.font.size=Pt(size); r.font.color.rgb=color
def bullet(t):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(t); r.font.size=Pt(10.5); r.font.color.rgb=DARK
def numstep(num,t):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.space_after=Pt(3)
    rn=p.add_run(f'{num}.  '); rn.bold=True; rn.font.size=Pt(10.5); rn.font.color.rgb=AMBER
    r=p.add_run(t); r.font.size=Pt(10.5); r.font.color.rgb=DARK
def callout(title,t,fill=AMBERBG,tc=AMBER):
    tb=doc.add_table(rows=1,cols=1); tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=tb.cell(0,0); shade(c,fill); c.text=''
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(title); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=tc
    p2=c.add_paragraph(); r2=p2.add_run(t); r2.font.size=Pt(10); r2.font.color.rgb=DARK
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
def table(headers,rows,widths=None):
    tb=doc.add_table(rows=1,cols=len(headers)); tb.alignment=WD_TABLE_ALIGNMENT.CENTER; tb.style='Table Grid'
    hd=tb.rows[0].cells
    for i,t in enumerate(headers): shade(hd[i],HEADBG); sct(hd[i],t,bold=True,white=True,size=9.5)
    for ri,row in enumerate(rows):
        cs=tb.add_row().cells
        for i,v in enumerate(row):
            sct(cs[i],str(v),size=9.5)
            if ri%2==1: shade(cs[i],ZEBRA)
    if widths:
        for i,w in enumerate(widths):
            for row in tb.rows: row.cells[i].width=Inches(w)

# ===== PORTADA =====
for _ in range(4): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('TOKENIA'); r.bold=True; r.font.size=Pt(28); r.font.color.rgb=AMBER
r2=p.add_run(' STUDIO'); r2.bold=True; r2.font.size=Pt(28); r2.font.color.rgb=DARK
for _ in range(2): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('MANUAL DE USUARIO'); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=DARK
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Módulo de Presupuesto · PyG Presupuesto vs Real'); r.font.size=Pt(13); r.font.color.rgb=AMBER
doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('FMV Dashboard — Cuadro de Mando Financiero'); r.font.size=Pt(12); r.font.color.rgb=GREY
for _ in range(7): doc.add_paragraph()
tb=doc.add_table(rows=0,cols=2); tb.alignment=WD_TABLE_ALIGNMENT.CENTER
for k,v in [
    ('Aplicación','FMV Dashboard'),
    ('Módulo','Presupuesto (Presupuesto vs Real)'),
    ('Dirigido a','Dirección y control financiero de FMV'),
    ('Versión','1.0'),
    ('Fecha','15 de junio de 2026'),
]:
    cs=tb.add_row().cells; cs[0].width=Inches(2.0); cs[1].width=Inches(4.0)
    sct(cs[0],k,bold=True,size=10,color=AMBER); sct(cs[1],v,size=10)
doc.add_page_break()

# ===== 0 INTRO =====
h1('Para qué sirve este módulo')
body('El módulo de Presupuesto compara, mes a mes y de forma acumulada, el PRESUPUESTO del ejercicio frente '
     'a los datos REALES de contabilidad, con la estructura de una cuenta de resultados analítica (PyG): '
     'ventas, márgenes, EBITDA y resultado. Permite ver de un vistazo dónde se está cumpliendo el '
     'presupuesto y dónde hay desviaciones, y bajar al detalle de cada cuenta y movimiento.')
h2('En una frase')
callout('Flujo de trabajo',
        '1) Cargas el diario contable (datos reales).  2) Cargas el fichero de presupuesto del año (GL PPT de '
        'Business Central).  3) El módulo te muestra Presupuesto vs Real con sus desviaciones, y puedes bajar '
        'al detalle y exportar.')

# ===== 1 REQUISITO PREVIO =====
h1('Requisito previo: datos contables', 1)
body('Para que el módulo funcione necesita PRIMERO los datos reales de contabilidad cargados en el Dashboard '
     '(el diario contable). Si no los hay, al entrar verás el mensaje:')
callout('"Sin datos contables"',
        'Carga un diario contable primero para ver el presupuesto vs real. Hasta entonces el módulo no puede '
        'comparar nada porque no tiene "lo real".', BLUEBG, DARK)
body('Una vez cargada la contabilidad, ya puedes cargar el presupuesto.')

# ===== 2 EL FICHERO =====
h1('El fichero de presupuesto', 2)
h2('Formato')
body('El formato recomendado y que el sistema reconoce de forma nativa es el GL PPT de Business Central: el '
     'informe "Presupuesto gastos e ingresos". El sistema lo detecta automáticamente porque las columnas de '
     'los meses vienen como fechas en la fila de cabecera.')
h2('Qué cuentas se tienen en cuenta')
bullet('Solo se procesan las cuentas de los grupos 6XX (gastos) y 7XX (ingresos).')
bullet('Los ingresos (7XX) vienen en negativo en el fichero; el sistema les ajusta el signo automáticamente.')
bullet('El máximo desglose es a 9 dígitos (subcuenta); el sistema agrega también a 3 dígitos y a categoría de PyG.')
h2('Formatos alternativos admitidos (si no usas el GL PPT)')
table(['Formato','Estructura'],
    [
        ['Horizontal','Una fila por cuenta y columnas Ene, Feb … Dic (o meses con nombre completo).'],
        ['Vertical','Columnas Cuenta, Mes, Importe (una fila por cuenta y mes).'],
    ],
    widths=[1.6,5.0])
callout('Importante sobre el año',
        'En el formato GL PPT el año va dentro del fichero (en las fechas de cabecera). En los formatos '
        'alternativos NO hay año en el fichero: el año lo fija siempre el desplegable "Año" al cargar.', AMBERBG, AMBER)

# ===== 3 CÓMO CARGAR =====
h1('Cómo cargar el presupuesto', 3)
body('En la pestaña Presupuesto, sigue estos pasos:')
numstep(1,'Localiza el bloque "Cargar Presupuesto" (si ya hay uno cargado, pulsa "Recargar presupuesto" para mostrarlo).')
numstep(2,'En el desplegable "Año", selecciona el ejercicio que vas a cargar (p. ej. 2027).')
numstep(3,'Pulsa "Seleccionar archivo" (o "Reemplazar" si ya había uno) y elige el Excel del presupuesto.')
numstep(4,'Confirma el ejercicio en el aviso que aparece (ver capítulo 4).')
numstep(5,'Espera al mensaje verde de confirmación: "Presupuesto [año] cargado: N líneas".')
callout('El año lo manda el desplegable',
        'El año con el que se guarda el presupuesto es el que eliges en el desplegable, NO el nombre del '
        'fichero. Antes de cargar 2027, asegúrate de que el desplegable dice 2027. El año 2027 ya aparece en '
        'la lista automáticamente.')
body('Nota: cargar un presupuesto REEMPLAZA por completo el que hubiera de ese mismo año. No se acumula.')

# ===== 4 PROTECCIONES =====
h1('Protecciones contra cargar en el ejercicio incorrecto', 4)
body('El sistema incorpora dos protecciones para evitar volcar, por ejemplo, el presupuesto de 2027 dentro '
     'de 2026 por dejar mal el desplegable:')
h2('Protección 1 — Confirmación del ejercicio')
body('Antes de cargar, aparece un aviso que te obliga a confirmar el año:')
callout('Aviso de confirmación',
        '"Vas a cargar el presupuesto en el ejercicio 2027. Si ya existe un presupuesto de 2027, se '
        'REEMPLAZARÁ por completo. ¿El año 2027 es correcto?"  →  Si pulsas Cancelar, no se carga nada.',
        BLUEBG, DARK)
h2('Protección 2 — Verificación automática del año (solo GL PPT)')
body('Como el GL PPT lleva el año real en las fechas de cabecera, el sistema lo lee y lo compara con el '
     'desplegable. Si no coinciden, BLOQUEA la carga y avisa:')
callout('Carga bloqueada',
        '"El fichero parece ser del año 2027, pero has seleccionado 2026. Cambia el año en el desplegable y '
        'vuelve a cargar."  →  No se guarda nada hasta que el año del desplegable coincida con el del fichero.',
        REDBG, RED)
body('Con estas dos protecciones, aunque te despistes con el selector, el dato del fichero manda y te frena '
     'antes de mezclar ejercicios.')

# ===== 5 LA TABLA =====
h1('La tabla Presupuesto vs Real', 5)
body('Una vez cargados contabilidad y presupuesto, verás la tabla "PyG Presupuesto vs Real [año]". En la '
     'cabecera tienes el selector de Mes y el botón Exportar.')
h2('Columnas')
table(['Columna','Significado'],
    [
        ['Concepto','Línea del PyG (Ventas, Compras, Margen Bruto, Personal, EBITDA, Resultado…).'],
        ['Ppto. Mes','Presupuesto del mes seleccionado.'],
        ['Real Mes','Dato real (contabilidad) del mes seleccionado.'],
        ['Var.','Desviación del mes en % (real frente a presupuesto).'],
        ['Ppto. Acum','Presupuesto acumulado de enero al mes seleccionado.'],
        ['Real Acum','Real acumulado de enero al mes seleccionado.'],
        ['Var. (acum.)','Desviación acumulada en %.'],
    ],
    widths=[1.6,5.0])
h2('Selector de mes')
body('El desplegable de la cabecera elige el mes. Las columnas "Mes" muestran ese mes; las columnas "Acum" '
     'muestran el acumulado desde enero hasta ese mes.')
h2('Estructura del PyG')
body('Las líneas siguen una cuenta de resultados analítica, con subtotales calculados automáticamente:')
bullet('VENTAS  (–) Compras  (±) Var. existencias  →  MARGEN BRUTO')
bullet('(–) Servicios ext.  (–) Personal  (+) Subvenciones / Otros ing. explotación  →  EBITDA')
bullet('(–) Resto gastos  (–) Amortizaciones  (–) Gastos financieros  (+) Ing. financieros / excepcionales  →  RESULTADO')
body('Las líneas opcionales (subvenciones, ingresos financieros, etc.) solo aparecen si tienen importe.')

# ===== 6 DESVIACIONES =====
h1('Interpretar las desviaciones (colores)', 6)
body('La columna de variación usa color según si la desviación es buena o mala, teniendo en cuenta si la '
     'línea es ingreso o gasto:')
table(['Tipo de línea','Verde (favorable)','Rojo (desfavorable)'],
    [
        ['Ingresos (ventas, etc.)','Real por encima del presupuesto','Real por debajo del presupuesto'],
        ['Gastos (compras, personal…)','Real por debajo del presupuesto','Real por encima del presupuesto'],
    ],
    widths=[2.2,2.2,2.2])
body('Es decir: el verde siempre significa "vamos mejor que lo presupuestado" y el rojo "peor", se trate de '
     'ingresos o de gastos.')

# ===== 7 DRILL-DOWN =====
h1('Bajar al detalle (drill-down)', 7)
body('Puedes desplegar cada línea para ver su composición, en tres niveles:')
numstep(1,'Clic en una categoría del PyG (con el triángulo ▶) → muestra las cuentas a 3 dígitos que la componen.')
numstep(2,'Clic en una cuenta de 3 dígitos → muestra sus subcuentas a 9 dígitos (si hay más de una).')
numstep(3,'Cada nivel muestra sus propias columnas de Ppto/Real/Variación, mes y acumulado.')
body('Vuelve a hacer clic para plegar. Los subtotales (Margen Bruto, EBITDA, Resultado) no se despliegan: son '
     'cálculos.')

# ===== 8 EXPORTAR =====
h1('Exportar a Excel', 8)
h2('Exportar toda la tabla')
body('El botón "Exportar" de la cabecera genera un Excel con todas las líneas visibles (Concepto, Ppto/Real '
     'del mes y acumulado, y variaciones). El fichero se nombra automáticamente con el año y el mes.')
h2('Exportar el detalle de movimientos')
body('Haciendo clic sobre cualquier importe de las columnas REAL (mes o acumulado) —en categoría, cuenta de 3 '
     'dígitos o subcuenta— se descarga un Excel con los MOVIMIENTOS contables reales que componen esa cifra '
     '(fecha, cuenta, descripción, debe, haber, neto, documento y proveedor). Ideal para justificar una '
     'desviación.')
callout('Truco',
        'Las cifras de la columna Real están subrayadas al pasar el ratón: ese subrayado indica que puedes '
        'pulsarlas para exportar su detalle de movimientos.')

# ===== 9 RECARGAR =====
h1('Recargar o cambiar el presupuesto', 9)
bullet('Para sustituir el presupuesto de un año: selecciona ese año, pulsa "Recargar presupuesto" y carga el '
       'fichero nuevo. Reemplaza por completo el anterior de ese año.')
bullet('Para cargar otro año (p. ej. añadir 2027 manteniendo 2026): cambia el desplegable de año y carga. '
       'Cada año se guarda por separado.')
bullet('El año que se muestra en la tabla es el ejercicio activo del Dashboard (selector de año general).')

# ===== 10 PROBLEMAS =====
h1('Problemas frecuentes', 10)
table(['Síntoma','Causa y solución'],
    [
        ['"Sin datos contables"','Falta cargar el diario contable. Cárgalo primero.'],
        ['"El fichero parece ser del año X…"','El año del desplegable no coincide con el del fichero. Ajusta el desplegable y recarga.'],
        ['"No se encontraron datos de presupuesto válidos"','El Excel no tiene cuentas 6XX/7XX reconocibles. Verifica que es el informe correcto (GL PPT).'],
        ['"El archivo está vacío"','La hoja no tiene datos o no es el fichero esperado.'],
        ['No aparece una línea (p. ej. subvenciones)','Es una línea opcional: solo se muestra si tiene importe en presupuesto o real.'],
    ],
    widths=[2.4,4.2])

doc.add_paragraph()
callout('Resumen rápido',
        'Carga contabilidad → selecciona el año correcto → carga el GL PPT y confirma el ejercicio → revisa '
        'desviaciones (verde = mejor, rojo = peor) → despliega para el detalle → exporta lo que necesites.',
        GREENBG, GREEN)

doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('— Fin del manual —'); r.font.size=Pt(10); r.font.color.rgb=LIGHT
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Powered by TOKENIA Studio'); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=AMBER

fp=doc.sections[0].footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fr=fp.add_run('TOKENIA Automation SLU · Manual Presupuesto · FMV Dashboard · v1.0 · 15/06/2026')
fr.font.size=Pt(8); fr.font.color.rgb=LIGHT

out=r'C:\TOKENIA-STUDIO\projects\FMV DASHBOARD\fmv-dashboard-v2\docs\Manual_Presupuesto_FMV.docx'
doc.save(out); print('OK ->', out)
